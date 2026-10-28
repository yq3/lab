"""TC-ING-01 多格式样例生成器（自造小型样例——不引外部下载；可重复执行，产物确定）。

产物（tests/fixtures/ 下）：
- sample.docx：多级标题（1/2 级）+ 表格 + 多段正文
- sample.pptx：两页 slide（标题+正文）
- sample.xlsx：两 sheet（税率表 3 列 3 行 / 免税政策 2 列 2 行）
- sample.html：h1/h2 两级标题 + 段落
- sample.pdf：原生 PDF（手写 PDF 字节——大字号标题可被 layout 模型识别为 section header）

运行：`uv run python tests/fixtures/gen_samples.py`（依赖 python-docx/python-pptx/openpyxl
——均为 docling 依赖，无新增）。
"""

from __future__ import annotations

from pathlib import Path

FIXTURES = Path(__file__).resolve().parent


def build_docx(path: Path) -> None:
    from docx import Document

    d = Document()
    d.add_heading("增值税申报指南", 1)
    d.add_heading("一般纳税人", 2)
    d.add_paragraph("一般纳税人销售自产货物并提供安装服务的，应分别核算货物的销售额和安装服务的销售额。")
    d.add_paragraph("一般纳税人发生应税销售行为，应当向索取增值税专用发票的购买方开具专用发票。")
    d.add_heading("小规模纳税人", 2)
    d.add_paragraph("小规模纳税人发生应税销售行为，适用简易计税方法计税。")
    t = d.add_table(rows=2, cols=3)
    t.style = "Table Grid"
    for r, vals in enumerate([("税种", "税率", "备注"), ("增值税", "13%", "标准税率")]):
        for c, v in enumerate(vals):
            t.cell(r, c).text = v
    d.save(path)


def build_pptx(path: Path) -> None:
    from pptx import Presentation

    p = Presentation()
    for title, body in (
        ("增值税要点", "进项税额凭合规扣税凭证抵扣。"),
        ("留抵退税", "符合条件的纳税人可申请退还增量留抵税额。"),
    ):
        slide = p.slides.add_slide(p.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    p.save(path)


def build_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "税率表"
    ws.append(["税种", "税率", "适用"])
    ws.append(["增值税", "13%", "一般货物"])
    ws.append(["增值税", "9%", "交通运输"])
    ws.append(["增值税", "6%", "现代服务"])
    ws2 = wb.create_sheet("免税政策")
    ws2.append(["项目", "条件"])
    ws2.append(["农产品", "自产自销"])
    ws2.append(["科技服务", "技术转让"])
    wb.save(path)


def build_html(path: Path) -> None:
    path.write_text(
        "<!doctype html><html><head><meta charset=\"utf-8\"><title>费用制度</title></head>"
        "<body><h1>费用制度总则</h1><p>差旅费报销标准如下，按城市类别分档执行。</p>"
        "<h2>住宿标准</h2><p>一类城市每晚不超过 600 元。</p></body></html>",
        encoding="utf-8",
    )


_PDF_LINES = [
    # (text, font_size, y)
    ("VAT Filing Guide 2026", 22, 740),
    ("Chapter 1 General Taxpayers", 16, 690),
    ("A general taxpayer shall account separately for sales of self-produced goods.", 11, 660),
    ("Chapter 2 Small-scale Taxpayers", 16, 620),
    ("The simplified taxation method applies to small-scale taxpayers.", 11, 590),
]


def build_pdf(path: Path) -> None:
    """最小原生 PDF（Helvetica 文本行；大字号行作标题——layout 模型可识别 section header）。"""
    objs: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
    ]
    streams = []
    for text, size, y in _PDF_LINES:
        esc = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        streams.append(f"BT /F1 {size} Tf 72 {y} Td ({esc}) Tj ET")
    stream = "\n".join(streams).encode()
    objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    ).encode()
    path.write_bytes(bytes(out))


def build_all(base: Path | None = None) -> None:
    base = base or FIXTURES
    build_docx(base / "sample.docx")
    build_pptx(base / "sample.pptx")
    build_xlsx(base / "sample.xlsx")
    build_html(base / "sample.html")
    build_pdf(base / "sample.pdf")
    print(f"samples written to {base}")


if __name__ == "__main__":
    build_all()
