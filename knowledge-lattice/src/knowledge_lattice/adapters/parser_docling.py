"""docling 封装（design §3 adapters/parser_docling.py——DoclingDocument → 统一块结构）。

统一块结构（TC-ING-01）：content / heading_path / page_no / bbox_json / sheet_name（+
kind=heading|text|table、heading level、表格 rows——结构边界保留供切分）。

- 标题层级：docling iterate_items 的嵌套层级（docx 根组占位使首项 level=2——按首项
  归一化 offset）；html 的 h1→TitleItem、h2+→SectionHeaderItem 同一归一化覆盖；
  PDF 标题经 layout 模型识别（大字号行），层级扁平（各自顶级章节）。
- 页码：PDF 逐项 prov.page_no；pptx=slide 序号；xlsx=sheet 序号；docx/html 无页码（None）。
- sheet_name：docling xlsx 后端不保留 sheet 名——经 openpyxl（docling 依赖，零新增）
  按 sheet 序号补名（FR-1）。
- TXT（TC-ING-04①，M1-4）：docling 无纯文本输入格式——本模块直读字节做**编码
  探测**（utf-8-sig → utf-8 → gb18030 依序严格解码，选正常字符比最高的候选；
  全部失败退 latin-1 由摄取侧 mojibake 质量门把关）→ 段落（空行分隔）为块。
- 老格式（TC-ING-04②）：`.doc/.xls/.ppt` 拒收（上传入口 400 提示转存新格式——
  零中间件纪律不引入 LibreOffice；此处纵深防御同拒）。

PDF 解析需 layout 模型权重（首运行下载；网络受限环境经 HF_ENDPOINT 镜像），
选项 do_ocr/do_table_structure 关闭以最小化模型面（M0a 无扫描件/无 PDF 表格样例）。
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

# 老格式拒收清单（TC-ING-04②——上传入口与服务层共用同一常量）
LEGACY_SUFFIXES = frozenset({".doc", ".xls", ".ppt"})

# TXT 编码探测候选（严格解码依序尝试；gb18030 为 GBK/GB2312 超集——TC-ING-04①）
TXT_ENCODING_CANDIDATES = ("utf-8-sig", "utf-8", "gb18030")


class UnsupportedFormatError(Exception):
    """老格式拒收（.doc/.xls/.ppt——提示转存新格式，TC-ING-04②）。"""


@dataclass(frozen=True)
class ParsedBlock:
    """统一块结构（parse 阶段产物，供 chunk 阶段消费）。"""

    kind: str  # heading | text | table
    content: str
    heading_path: tuple[str, ...] = ()
    level: int | None = None  # heading 层级（1 起）
    page_no: int | None = None
    bbox_json: str | None = None
    sheet_name: str | None = None
    rows: tuple[tuple[str, ...], ...] | None = None  # 表格行（整块不拆 / Excel 行块切分）


@dataclass(frozen=True)
class ParsedDocument:
    blocks: tuple[ParsedBlock, ...]
    page_count: int | None = None
    parser: str = "docling"


@lru_cache(maxsize=1)
def _converter():
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption

    pdf_opts = PdfPipelineOptions()
    pdf_opts.do_ocr = False
    pdf_opts.do_table_structure = False
    return DocumentConverter(
        format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=pdf_opts)}
    )


def _bbox_json(prov) -> str | None:
    bbox = getattr(prov, "bbox", None)
    if bbox is None:
        return None
    return json.dumps(
        {"l": round(bbox.l, 2), "t": round(bbox.t, 2), "r": round(bbox.r, 2), "b": round(bbox.b, 2)}
    )


def _rows_from_table_cells(table_cells) -> tuple[tuple[str, ...], ...]:
    """docling TableCell（含 start_row/col_offset_idx 与 col_span）→ 规整行网格。"""
    if not table_cells:
        return (("",),)
    n_rows = max(c.end_row_offset_idx for c in table_cells)
    n_cols = max(c.end_col_offset_idx for c in table_cells)
    grid = [[""] * n_cols for _ in range(n_rows)]
    for c in table_cells:
        for r in range(c.start_row_offset_idx, c.end_row_offset_idx):
            for col in range(c.start_col_offset_idx, c.end_col_offset_idx):
                grid[r][col] = c.text
    return tuple(tuple(row) for row in grid)


def _table_content(rows: tuple[tuple[str, ...], ...]) -> str:
    return "\n".join(" | ".join(row) for row in rows)


def _sheet_names(path: Path) -> list[str]:
    """xlsx sheet 名清单（按 sheet 序号——docling TableItem 的 page_no 对位）。"""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def decode_txt_bytes(data: bytes) -> tuple[str, str]:
    """TXT 编码探测（TC-ING-04①）：候选严格解码 → 取正常字符比最高者
    （「探测并转 UTF-8」——解码即转；全部失败退 latin-1，质量由 mojibake 门把关）。

    返回 (text, encoding)。
    """
    from knowledge_lattice.core.docs.quality import normal_char_ratio

    best: tuple[str, str] | None = None
    best_ratio = -1.0
    for enc in TXT_ENCODING_CANDIDATES:
        try:
            text = data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
        ratio = normal_char_ratio(text)
        if ratio >= 0.99:
            return text, enc  # 干净解码——无需再比
        if ratio > best_ratio:
            best, best_ratio = (text, enc), ratio
    if best is not None:
        return best
    return data.decode("latin-1"), "latin-1"


def _parse_txt(path: Path) -> ParsedDocument:
    """TXT → 统一块结构（空行分段；无页码/表格——纯文本形态）。"""
    data = path.read_bytes()
    text, encoding = decode_txt_bytes(data)
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    blocks = tuple(
        ParsedBlock(kind="text", content=p) for p in paragraphs
    )
    doc = ParsedDocument(blocks=blocks, parser=f"txt:{encoding}")
    return doc


def parse_document(path: str | Path) -> ParsedDocument:
    """解析任意支持格式 → 统一块结构（docling 原生：docx/pptx/xlsx/html/pdf；
    TXT 本模块编码探测；.doc/.xls/.ppt 拒收——UnsupportedFormatError）。"""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in LEGACY_SUFFIXES:
        raise UnsupportedFormatError(
            f"老格式 {suffix} 不受支持——请转存为新格式（.docx/.xlsx/.pptx）后重传"
            "（零中间件纪律：不引入 LibreOffice，TC-ING-04②）"
        )
    if suffix == ".txt":
        return _parse_txt(path)
    conv = _converter()
    result = conv.convert(str(path))
    doc = result.document

    sheet_names = _sheet_names(path) if path.suffix.lower() == ".xlsx" else []
    blocks: list[ParsedBlock] = []
    heading_stack: list[str] = []
    first_level: int | None = None
    last_page: int | None = None

    for item, level in doc.iterate_items():
        if first_level is None:
            first_level = level
        offset = first_level - 1  # docx 根组占位归一化（html/pdf 首项 level=1 → offset=0）
        item_type = type(item).__name__
        prov = getattr(item, "prov", None) or []
        page_no = prov[0].page_no if prov else None
        bbox_json = _bbox_json(prov[0]) if prov else None
        text = getattr(item, "text", None)

        if item_type == "TableItem":
            rows = _rows_from_table_cells(getattr(item.data, "table_cells", None) or [])
            sheet = sheet_names[page_no - 1] if sheet_names and page_no and 0 < page_no <= len(sheet_names) else None
            blocks.append(
                ParsedBlock(
                    kind="table",
                    content=_table_content(rows),
                    heading_path=tuple(heading_stack),
                    page_no=page_no,
                    bbox_json=bbox_json,
                    sheet_name=sheet,
                    rows=rows,
                )
            )
            continue

        label = str(getattr(item, "label", ""))
        is_heading = item_type in ("SectionHeaderItem", "TitleItem")
        content = (text or "").strip()
        if not content:
            continue

        if is_heading:
            lvl = max(1, level - offset)
            if item_type == "TitleItem" and lvl > 1:
                lvl = 1  # 文档主标题（pptx 每页标题 / html h1）
            # pptx：翻页重置标题栈（每页标题独立）
            if last_page is not None and page_no is not None and page_no != last_page:
                heading_stack.clear()
            last_page = page_no
            del heading_stack[lvl - 1:]
            heading_stack.append(content)
            blocks.append(
                ParsedBlock(
                    kind="heading",
                    content=content,
                    heading_path=tuple(heading_stack[:-1]),
                    level=lvl,
                    page_no=page_no,
                    bbox_json=bbox_json,
                )
            )
        else:
            blocks.append(
                ParsedBlock(
                    kind="text",
                    content=content,
                    heading_path=tuple(heading_stack),
                    page_no=page_no,
                    bbox_json=bbox_json,
                )
            )

    try:
        page_count = doc.num_pages()
    except Exception:
        page_count = None
    return ParsedDocument(blocks=tuple(blocks), page_count=page_count)
